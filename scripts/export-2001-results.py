#!/usr/bin/env python3
"""Export the 2001 worksheet from the Commons Library historical-results workbook."""
import argparse
import gzip
import hashlib
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = ROOT / "data/source/hoc/1918-2019election_results_by_pcon.xlsx"
DEFAULT_BOUNDARIES = ROOT / "data/source/boundaries/constituencies-2001.geojson.gz"
DEFAULT_OUTPUT = ROOT / "data/source/hoc/historical-2001-results.json"
OFFICIAL_TOTALS = {"seats": 659, "electorate": 44403238, "validVotes": 26367383}
EXPECTED_COUNTRIES = {"England": 529, "Wales": 40, "Scotland": 72, "Northern Ireland": 18}
PARTY_COLUMNS = [
    ("Conservative", "Conservative", 8),
    ("Liberal Democrats", "Liberal Democrats", 11),
    ("Labour", "Labour", 14),
    ("SNP", "Scottish National", 17),
    ("Plaid Cymru", "Plaid Cymru", 20),
    ("DUP", "Democratic Unionist", 23),
    ("Sinn Fein", "Sinn Féin", 26),
    ("SDLP", "Social Democratic and Labour", 29),
    ("UUP", "Ulster Unionist", 32),
    ("Other", "Other", 35),
]
NAME_ALIASES = {
    "chester city of": "city of chester",
    "durham city of": "city of durham",
    "regent s park and kensington north": "regent s park and north kensington",
    "richmond": "richmond yorks",
    "wrekin the": "the wrekin",
    "york city of": "city of york",
}
OTHER_WINNER_NOTES = {
    "glasgow springburn": {
        "party": "Speaker",
        "partyOfficial": "Speaker",
        "member": "Michael Martin",
        "votes": 16053,
        "note": "The Commons Library worksheet notes that Michael Martin won standing as Speaker.",
    },
    "wyre forest": {
        "party": "Independent",
        "partyOfficial": "Independent Kidderminster Hospital and Health Concern",
        "member": "Richard Taylor",
        "votes": 28487,
        "note": "The Commons Library worksheet notes that Richard Taylor won standing as Independent Kidderminster Hospital and Health Concern.",
    },
}

def normalized(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii").lower()
    text = text.replace("&", " and ")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())

def digest(path):
    hasher = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()

def integer(value, label):
    if value is None:
        return 0
    number = float(value)
    if not number.is_integer() or number < 0:
        raise ValueError(f"{label} is not a non-negative whole number: {value!r}")
    return int(number)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--boundaries", type=Path, default=DEFAULT_BOUNDARIES)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    boundary_path = args.boundaries.resolve()
    boundaries = json.loads(gzip.open(boundary_path, "rt", encoding="utf-8").read())
    boundary_by_name = {normalized(feature["properties"]["name"]): feature for feature in boundaries["features"]}
    if len(boundary_by_name) != 659:
        raise ValueError(f"Expected 659 unique 2001 boundary names; found {len(boundary_by_name)}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "2001" not in workbook.sheetnames:
        raise ValueError("Workbook has no 2001 worksheet")
    sheet = workbook["2001"]
    if sheet["C4"].value != "Constituency" or sheet["AM4"].value != "Total votes":
        raise ValueError("Unexpected 2001 worksheet layout")

    seats = []
    matched_ids = set()
    country_counts = Counter()
    electorate_total = 0
    valid_total = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=5, max_row=663, values_only=True), start=5):
        if row[1] is None:
            continue
        workbook_id = integer(row[1], f"row {row_number} workbook id")
        workbook_name = str(row[2]).strip()
        name_key = normalized(workbook_name)
        boundary_key = NAME_ALIASES.get(name_key, name_key)
        feature = boundary_by_name.get(boundary_key)
        if feature is None:
            raise ValueError(f"Cannot match 2001 result row {row_number}: {workbook_name}")
        props = feature["properties"]
        if props["id"] in matched_ids:
            raise ValueError(f"Duplicate boundary match for {workbook_name}: {props['id']}")
        matched_ids.add(props["id"])

        electorate = integer(row[6], f"{workbook_name} electorate")
        valid_votes = integer(row[38], f"{workbook_name} valid votes")
        if electorate <= 0 or valid_votes <= 0 or valid_votes > electorate:
            raise ValueError(f"Invalid electorate/valid votes for {workbook_name}")
        party_totals = []
        for official, party, column in PARTY_COLUMNS:
            votes = integer(row[column], f"{workbook_name} {official} votes")
            if votes:
                party_totals.append({
                    "party": party,
                    "partyOfficial": official,
                    "votes": votes,
                    "share": votes / valid_votes,
                    "isGrouped": official == "Other",
                })
        if sum(item["votes"] for item in party_totals) != valid_votes:
            raise ValueError(f"Party groups do not sum to valid votes in {workbook_name}")

        top_votes = max(item["votes"] for item in party_totals)
        top_rows = [item for item in party_totals if item["votes"] == top_votes]
        if len(top_rows) != 1:
            raise ValueError(f"Ambiguous leading party group in {workbook_name}")
        top = top_rows[0]
        note = None
        member = None
        winner_votes = top_votes
        if top["party"] == "Other":
            override = OTHER_WINNER_NOTES.get(name_key)
            if not override:
                raise ValueError(f"Grouped Other leads in {workbook_name} without a source note")
            party = override["party"]
            party_official = override["partyOfficial"]
            member = override["member"]
            winner_votes = override["votes"]
            note = override["note"]
        else:
            party = top["party"]
            party_official = top["partyOfficial"]

        country = str(row[5]).strip()
        region = str(row[4]).strip() if row[4] else country
        source_turnout = float(row[39]) if row[39] is not None else None
        turnout = valid_votes / electorate
        if source_turnout is not None and abs(source_turnout - turnout) > 0.002:
            raise ValueError(f"Source turnout disagrees with valid votes/electorate in {workbook_name}")
        seats.append({
            "id": props["id"],
            "name": props["name"],
            "country": country,
            "region": region,
            "county": str(row[3]).strip() if row[3] else None,
            "type": "constituency",
            "gss": None,
            "sourceCode": props.get("sourceCode"),
            "sourceWorkbookId": workbook_id,
            "electorate": electorate,
            "validVotes": valid_votes,
            "invalidVotes": None,
            "invalidVotesKnown": False,
            "turnout": turnout,
            "party": party,
            "partyGroup": party,
            "partyOfficial": party_official,
            "firstParty": party,
            "secondParty": None,
            "winnerVotes": winner_votes,
            "winnerShare": winner_votes / valid_votes,
            "majority": None,
            "majorityShare": None,
            "resultType": None,
            "gainedFrom": None,
            "member": member,
            "memberGender": None,
            "candidates": [],
            "partyTotals": party_totals,
            "notes": note,
        })
        country_counts[country] += 1
        electorate_total += electorate
        valid_total += valid_votes

    if len(seats) != OFFICIAL_TOTALS["seats"] or matched_ids != {f["properties"]["id"] for f in boundaries["features"]}:
        raise ValueError(f"Expected an exact 659-seat result/boundary match; found {len(seats)} results and {len(matched_ids)} matches")
    if dict(country_counts) != EXPECTED_COUNTRIES:
        raise ValueError(f"Unexpected country seat coverage: {dict(country_counts)}")
    if electorate_total != OFFICIAL_TOTALS["electorate"] or valid_total != OFFICIAL_TOTALS["validVotes"]:
        raise ValueError(f"National totals differ from reviewed figures: electorate={electorate_total}, valid={valid_total}")
    if sum(item["votes"] for seat in seats for item in seat["partyTotals"]) != valid_total:
        raise ValueError("National party-group votes do not sum to valid votes")

    workbook_hash = digest(workbook_path)
    election = {
        "id": "2001",
        "year": 2001,
        "date": "2001-06-07",
        "label": "2001",
        "boundarySetId": "2001",
        "resultCodePeriod": "1997–2001 PCA identifiers",
        "isNotional": False,
        "expectedSeatCount": 659,
        "marginDataAvailable": False,
        "candidateDataGranularity": "party-aggregate",
        "sourceId": "commons-library-2001-results",
        "sourceName": "House of Commons Library, CBP-8647 (2001 workbook tab)",
        "sourceUrl": "https://commonslibrary.parliament.uk/research-briefings/cbp-8647/",
        "sourceLicense": "Open Parliament Licence v3.0",
        "boundaryLabel": "2001 election boundaries",
        "caveat": "The workbook reports party-group totals, not named candidates. It does not provide constituency winning margins or invalid-ballot counts. “Other” combines smaller parties; the two seats led by that group are assigned from the workbook’s own notes.",
        "boundaryCaveat": "This map uses 2001 geometry. The Commons Library groups 1997–2001 PCA result identifiers, but reports minor interim changes in London and South East before 2001; do not treat this geometry as exact for 1997. Northern Ireland polygons union 1993 OSNI wards using the 1995 Order, which refers to ward areas as at 1 June 1994, leaving a one-year vintage mismatch at 1:50,000 source detail.",
        "seats": seats,
    }
    payload = {
        "schemaVersion": 1,
        "source": {
            "id": "commons-library-2001-results",
            "name": "House of Commons Library, CBP-8647",
            "url": "https://commonslibrary.parliament.uk/research-briefings/cbp-8647/",
            "license": "Open Parliament Licence v3.0",
            "attribution": "House of Commons Library",
            "inputFile": "1918-2019election_results_by_pcon.xlsx, 2001 worksheet",
            "inputOrigin": "Local GitHub mirror; not byte-compared with the Commons Library download",
            "inputSha256": workbook_hash,
            "validation": "659 unique rows, exact boundary-name join, official seat/electorate/valid-vote totals, and per-seat party-group arithmetic",
            "limits": "Party-group results; no candidate list, individual second-place votes, majority, or invalid-ballot counts",
        },
        "elections": [election],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + chr(10), encoding="utf-8")
    print(f"Exported {len(seats)} seats; electorate {electorate_total}; valid votes {valid_total}.")
    print(f"Workbook SHA-256: {workbook_hash}")
    print(f"Output: {args.output}")

if __name__ == "__main__":
    main()
